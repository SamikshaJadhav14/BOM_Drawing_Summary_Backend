from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def _write_table(worksheet, title, headers, rows, header_color):
    worksheet.append([title])
    worksheet.cell(worksheet.max_row, 1).font = Font(bold=True, size=14)
    worksheet.append(headers)

    header_row = worksheet.max_row
    for cell in worksheet[header_row]:
        cell.fill = PatternFill("solid", fgColor=header_color)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in rows:
        worksheet.append(row)

    for row in worksheet.iter_rows(
        min_row=header_row,
        max_row=worksheet.max_row,
        min_col=1,
        max_col=len(headers),
    ):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    worksheet.append([])


def generate_abstract_excel(summary, output_path):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "BOM Summary"

    drawing = summary.get("drawing_details", {})
    worksheet.append(["Drawing Details"])
    worksheet.cell(worksheet.max_row, 1).font = Font(bold=True, size=14)

    for key, value in drawing.items():
        worksheet.append([str(key).replace("_", " ").title(), value])

    worksheet.append([])

    abstract_rows = summary.get("abstract", [])
    _write_table(
        worksheet,
        "Abstract",
        ["SR. NO.", "SECTION", "TOTAL"],
        [
            [
                row.get("sr_no", ""),
                row.get("section", ""),
                row.get("total", 0),
            ]
            for row in abstract_rows
        ],
        "1F4E78",
    )

    parts = summary.get("parts", [])
    _write_table(
        worksheet,
        "Parts Summary",
        ["Part No.", "Description", "Specification", "Size", "Qty", "Weight (kg)"],
        [
            [
                part.get("part_no", ""),
                part.get("description", ""),
                part.get("specification", ""),
                part.get("size", ""),
                part.get("quantity", ""),
                part.get("weight_kg", ""),
            ]
            for part in parts
        ],
        "7030A0",
    )

    hardware = summary.get("hardware", [])
    _write_table(
        worksheet,
        "Hardware Summary",
        ["Description", "Purchase Item Code", "Qty", "Weight (kg)", "UOM"],
        [
            [
                item.get("description", ""),
                item.get("purchase_item_code", ""),
                item.get("quantity", ""),
                item.get("weight_kg", ""),
                item.get("uom", ""),
            ]
            for item in hardware
        ],
        "C65911",
    )

    dimensions = summary.get("dimension_summary", [])
    _write_table(
        worksheet,
        "Common Dimension Summary",
        ["Combined Size (LxWxT)", "Specification", "Part Count", "Part Nos."],
        [
            [
                row.get("combined_size", ""),
                row.get("specification", ""),
                row.get("part_count", ""),
                row.get("part_nos", ""),
            ]
            for row in dimensions
        ],
        "548235",
    )

    for column_cells in worksheet.columns:
        column_letter = get_column_letter(column_cells[0].column)
        width = max(len(str(cell.value or "")) for cell in column_cells)
        worksheet.column_dimensions[column_letter].width = min(max(width + 2, 12), 45)

    worksheet.freeze_panes = "A2"
    workbook.save(output_path)
